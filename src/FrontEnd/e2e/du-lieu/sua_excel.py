"""Sửa file Excel nhập liệu cho e2e "nhập lại, không build lại" (tester). Chạy: python -X utf8 -I sua_excel.py <thư-mục-excel> <lệnh>.

Lệnh: bo-dong-thuong-gap | doi-thu-tu | them-thu-tuc-xa | them-o | them-thu-bay | bo-gio-thu-tu | gio-hong | ma-la-thuong-gap
Chỉ dùng cho dữ liệu thử local (thư mục sinh bằng DichVuCong.DuLieuMau từ data-mau/). Cần openpyxl.
"""

import sys
from pathlib import Path

import openpyxl


def thu_tuc(thu_muc: Path):
    return thu_muc / "thu-tuc.xlsx"


def bo_dong_thuong_gap(thu_muc: Path) -> None:
    wb = openpyxl.load_workbook(thu_tuc(thu_muc))
    ws = wb["Thường gặp"]
    hang = [(r[0].value, r[1].value) for r in ws.iter_rows(min_row=2) if r[0].value]
    hang = [h for h in hang if h[1] != 14]
    ws.delete_rows(2, ws.max_row)
    for i, (ma, _) in enumerate(hang, start=1):
        ws.cell(row=i + 1, column=1, value=ma)
        ws.cell(row=i + 1, column=2, value=i)
    wb.save(thu_tuc(thu_muc))


def doi_thu_tu(thu_muc: Path) -> None:
    wb = openpyxl.load_workbook(thu_tuc(thu_muc))
    ws = wb["Thường gặp"]
    a, b = ws.cell(row=2, column=1), ws.cell(row=3, column=1)
    a.value, b.value = b.value, a.value
    wb.save(thu_tuc(thu_muc))


def them_thu_tuc_xa(thu_muc: Path) -> None:
    wb = openpyxl.load_workbook(thu_tuc(thu_muc))
    ws = wb["Thủ tục"]
    mau = [c.value for c in ws[2]]
    moi = list(mau)
    moi[0] = "T10.001"
    moi[1] = "Thủ tục thêm bởi tester T10"
    moi[3] = "x"  # Danh mục cấp xã
    moi[4] = None  # không thuộc cấp tỉnh
    ws.append(moi)
    wb.save(thu_tuc(thu_muc))


def them_o(thu_muc: Path) -> None:
    wb = openpyxl.load_workbook(thu_muc / "trung-tam.xlsx")
    wb["Ô dịch vụ"].append(["Ô mới T10", None, "lien-ket", "https://moi.example.test"])
    wb.save(thu_muc / "trung-tam.xlsx")


def them_thu_bay(thu_muc: Path) -> None:
    wb = openpyxl.load_workbook(thu_muc / "trung-tam.xlsx")
    ws = wb["Giờ làm việc"]
    # File mẫu đã có đủ 7 dòng thứ (ngày chưa điền giờ = nghỉ): điền giờ cho dòng Thứ Bảy.
    for r in ws.iter_rows(min_row=2):
        if r[0].value == "Thứ Bảy":
            r[1].value, r[2].value, r[3].value, r[4].value = "07:00", "11:30", None, None
    wb.save(thu_muc / "trung-tam.xlsx")


def bo_gio_thu_tu(thu_muc: Path) -> None:
    """Xoá giờ Thứ Tư (ngày nghỉ chen giữa) -> khối giờ thành hai nhóm."""
    wb = openpyxl.load_workbook(thu_muc / "trung-tam.xlsx")
    for r in wb["Giờ làm việc"].iter_rows(min_row=2):
        if r[0].value == "Thứ Tư":
            for c in r[1:5]:
                c.value = None
    wb.save(thu_muc / "trung-tam.xlsx")


def gio_hong(thu_muc: Path) -> None:
    """Giờ không hợp lệ: Thứ Hai sáng từ 25:00."""
    wb = openpyxl.load_workbook(thu_muc / "trung-tam.xlsx")
    for r in wb["Giờ làm việc"].iter_rows(min_row=2):
        if r[0].value == "Thứ Hai":
            r[1].value = "25:00"
    wb.save(thu_muc / "trung-tam.xlsx")


def ma_la_thuong_gap(thu_muc: Path) -> None:
    """Dòng Thường gặp trỏ tới mã không có trong sheet Thủ tục."""
    wb = openpyxl.load_workbook(thu_tuc(thu_muc))
    ws = wb["Thường gặp"]
    ws.append(["ZZ.999", 17])
    wb.save(thu_tuc(thu_muc))


LENH = {
    "bo-dong-thuong-gap": bo_dong_thuong_gap,
    "doi-thu-tu": doi_thu_tu,
    "them-thu-tuc-xa": them_thu_tuc_xa,
    "them-o": them_o,
    "them-thu-bay": them_thu_bay,
    "bo-gio-thu-tu": bo_gio_thu_tu,
    "gio-hong": gio_hong,
    "ma-la-thuong-gap": ma_la_thuong_gap,
}

if __name__ == "__main__":
    if len(sys.argv) != 3 or sys.argv[2] not in LENH:
        sys.exit(f"Cách dùng: sua_excel.py <thư-mục> <{'|'.join(LENH)}>")
    LENH[sys.argv[2]](Path(sys.argv[1]))
